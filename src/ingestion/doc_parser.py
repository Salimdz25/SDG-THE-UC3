"""
Module d'Extraction Documentaire Multi-formats pour l'UC3
Prend en charge : PDF, Word (.docx), Excel (.xlsx, .xls), CSV
Gère l'extraction textuelle, la pagination et les métadonnées.
"""

from pathlib import Path
from typing import Dict, Any, List
import pandas as pd
from pypdf import PdfReader
import docx

class DocumentParser:
    def __init__(self):
        pass

    def parse_file(self, file_path: str | Path) -> Dict[str, Any]:
        path = Path(file_path)
        if not path.exists():
            raise FileNotFoundError(f"Fichier introuvable: {path}")

        ext = path.suffix.lower()
        if ext == ".pdf":
            return self.parse_pdf(path)
        elif ext in (".docx", ".doc"):
            return self.parse_docx(path)
        elif ext in (".xlsx", ".xls"):
            return self.parse_excel(path)
        elif ext == ".csv":
            return self.parse_csv(path)
        elif ext in (".txt", ".md"):
            return self.parse_text(path)
        else:
            raise ValueError(f"Format non supporté: {ext}")

    def parse_pdf(self, path: Path) -> Dict[str, Any]:
        reader = PdfReader(str(path))
        pages_content = []
        full_text = []

        for idx, page in enumerate(reader.pages):
            text = page.extract_text() or ""
            pages_content.append({"page_number": idx + 1, "text": text.strip()})
            full_text.append(text.strip())

        return {
            "source_type": "PDF",
            "file_name": path.name,
            "file_path": str(path),
            "total_pages": len(reader.pages),
            "pages": pages_content,
            "full_text": "\n\n".join(full_text),
            "is_multipage": len(reader.pages) > 1
        }

    def parse_docx(self, path: Path) -> Dict[str, Any]:
        doc = docx.Document(str(path))
        paragraphs = [p.text for p in doc.paragraphs if p.text.strip()]
        
        # Extraire aussi le texte des tableaux s'il y en a
        tables_text = []
        for table in doc.tables:
            for row in table.rows:
                row_text = [cell.text.strip() for cell in row.cells if cell.text.strip()]
                if row_text:
                    tables_text.append(" | ".join(row_text))

        combined_text = "\n".join(paragraphs)
        if tables_text:
            combined_text += "\n\n--- Tableaux extraits ---\n" + "\n".join(tables_text)

        return {
            "source_type": "DOCX",
            "file_name": path.name,
            "file_path": str(path),
            "total_paragraphs": len(paragraphs),
            "full_text": combined_text,
            "is_multipage": False
        }

    def parse_excel(self, path: Path) -> Dict[str, Any]:
        excel_file = pd.ExcelFile(path)
        sheets_data = {}
        text_summary = []

        for sheet_name in excel_file.sheet_names:
            df = pd.read_excel(excel_file, sheet_name=sheet_name)
            sheets_data[sheet_name] = df.to_dict(orient="records")
            text_summary.append(f"Feuille: {sheet_name}\n" + df.to_string(index=False))

        return {
            "source_type": "EXCEL",
            "file_name": path.name,
            "file_path": str(path),
            "sheet_names": excel_file.sheet_names,
            "full_text": "\n\n".join(text_summary),
            "raw_data": sheets_data
        }

    def parse_csv(self, path: Path) -> Dict[str, Any]:
        df = pd.read_csv(path)
        return {
            "source_type": "CSV",
            "file_name": path.name,
            "file_path": str(path),
            "full_text": df.to_string(index=False),
            "raw_data": df.to_dict(orient="records")
        }

    def parse_text(self, path: Path) -> Dict[str, Any]:
        text = path.read_text(encoding="utf-8", errors="ignore")
        return {
            "source_type": "TEXT",
            "file_name": path.name,
            "file_path": str(path),
            "full_text": text
        }
