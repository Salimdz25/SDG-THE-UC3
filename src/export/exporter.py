"""
Générateur d'Exports Officiels Excel & Word pour la Validation Humaine UC3
Prêt pour soumission et archivage officiel au niveau du Rectorat et du portail THE.
"""

from pathlib import Path
from typing import List, Dict, Any
import pandas as pd
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

class UC3ReportExporter:
    def __init__(self):
        pass

    def export_excel(self, fiches_data: List[Dict[str, Any]], output_path: str | Path) -> str:
        """
        Génère un classeur Excel complet avec toutes les colonnes requises par l'audit UC3.
        """
        rows = []
        for f in fiches_data:
            rows.append({
                "ODD & Indicateur": f.get("odd_indicator", ""),
                "Titre Indicateur": f.get("indicator_title", ""),
                "Exigence THE 2027": f.get("methodological_requirement", ""),
                "Information Trouvée": f.get("information_found", ""),
                "Année": f.get("year", ""),
                "Source Exacte": f.get("source_exact", ""),
                "Citation Justificative": f.get("justifying_quote", ""),
                "Entité UC3": f.get("uc3_entity", ""),
                "Qualité Preuve": f.get("quality", "").upper(),
                "Caractère Public": f.get("publicity", "").upper(),
                "Niveau de Confiance": f.get("confidence", "").upper(),
                "Catégorie de Fiabilité": f.get("status_category", ""),
                "Points THE": f.get("the_points", 0.0),
                "Max Points": f.get("the_max_points", 3.0),
                "Score %": f.get("the_percentage", 0.0),
                "Alertes & Non-conformités": f.get("gap_or_alert", ""),
                "Action Proposée": f.get("proposed_action", "").upper(),
                "Validation Humaine": f.get("human_validation_status", "En attente"),
                "Responsable": f.get("assigned_responsible", ""),
                "Échéance": f.get("due_date", "")
            })

        df = pd.DataFrame(rows)
        out = Path(output_path)
        out.parent.mkdir(parents=True, exist_ok=True)
        df.to_excel(out, index=False, sheet_name="Fiches Indicateurs UC3")
        return str(out)

    def export_word(self, pilot_data: Dict[str, Any], output_path: str | Path) -> str:
        """
        Génère un rapport Word officiel (.docx) avec charte institutionnelle UC3.
        """
        doc = Document()
        
        # Titre Principal
        title_p = doc.add_paragraph()
        title_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        title_run = title_p.add_run("UNIVERSITÉ CONSTANTINE 3 - SALAH BOUBNIDER\n")
        title_run.bold = True
        title_run.font.size = Pt(16)
        title_run.font.color.rgb = RGBColor(0, 51, 102)

        sub_run = title_p.add_run("DOSSIER D'ÉVALUATION ET PREUVES DE DURABILITÉ - THE 2027\n")
        sub_run.bold = True
        sub_run.font.size = Pt(13)
        sub_run.font.color.rgb = RGBColor(180, 40, 40)

        meta_p = doc.add_paragraph(f"Année cible : {pilot_data.get('target_year', 2025)} | ODD Pilote : {pilot_data.get('sdg_name', 'ODD 17')}")
        meta_p.alignment = WD_ALIGN_PARAGRAPH.CENTER

        doc.add_heading("1. Synthèse de la Collecte et Score Prévisionnel", level=1)
        p_stats = doc.add_paragraph()
        p_stats.add_run(f"• Nombre total d'indicateurs analysés : {pilot_data.get('total_indicators', 0)}\n")
        p_stats.add_run(f"• Points THE estimés : {pilot_data.get('total_points_earned', 0)} / {pilot_data.get('max_possible_points', 0)} ")
        p_stats.add_run(f"({pilot_data.get('sdg_completion_percentage', 0)}%)\n")

        status_summary = pilot_data.get("status_summary", {})
        p_stats.add_run("• Répartition de la fiabilité des données :\n")
        for cat, cnt in status_summary.items():
            p_stats.add_run(f"   - {cat} : {cnt}\n")

        doc.add_heading("2. Fiches Détaillées par Indicateur", level=1)

        for f in pilot_data.get("fiches", []):
            doc.add_heading(f"Indicateur {f.get('odd_indicator')} : {f.get('indicator_title')}", level=2)
            
            table = doc.add_table(rows=0, cols=2)
            table.style = "Table Grid"
            
            fields = [
                ("Exigence méthodologique THE", f.get("methodological_requirement")),
                ("Information trouvée", f.get("information_found")),
                ("Année concernée", str(f.get("year", "N/A"))),
                ("Source exacte", f.get("source_exact")),
                ("Extrait justificatif (verbatim)", f.get("justifying_quote")),
                ("Entité UC3 concernée", f.get("uc3_entity")),
                ("Qualité de la preuve", f.get("quality").upper()),
                ("Caractère public", f.get("publicity").upper()),
                ("Niveau de confiance", f.get("confidence").upper()),
                ("Catégorie de fiabilité", f.get("status_category")),
                ("Score THE estimé", f"{f.get('the_points')} / {f.get('the_max_points')} pts ({f.get('the_percentage')}%)"),
                ("Lacune ou Alerte THE", f.get("gap_or_alert")),
                ("Action proposée", f.get("proposed_action").upper()),
                ("Statut de validation", f.get("human_validation_status")),
                ("Responsable assigné", f.get("assigned_responsible"))
            ]

            for label, val in fields:
                row = table.add_row()
                row.cells[0].paragraphs[0].add_run(label).bold = True
                row.cells[0].width = Inches(2.2)
                row.cells[1].paragraphs[0].add_run(str(val or ""))
                row.cells[1].width = Inches(4.3)

            doc.add_paragraph("")  # Espace

        out = Path(output_path)
        out.parent.mkdir(parents=True, exist_ok=True)
        doc.save(str(out))
        return str(out)
